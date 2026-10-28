"""Actualización de pesos de 2016-17 a pesos de 2025 con el IPC (WP6, PD5).

Qué hace: para cada registro de gasto de ENPH multiplica el valor por ``índice_base / índice_referencia``. «Deflactar a 2025» es en
rigor ACTUALIZAR a pesos de 2025 (llevar hacia adelante): el factor es > 1.

Decisiones de la autora (27 oct 2026)
  * Deflactor ÚNICO: el índice TOTAL del IPC nacional (hoja ``IndicesIPC``, serie de empalme 2003-2026 del DANE, base dic-2018 = 100;
    archivo ``anex-IPC-Indices-sep2026.xlsx``). NO es por división COICOP ni por ciudad: no se dispone de la serie por división para
    todo el periodo de la ENPH. Consecuencia: la actualización no cambia los pesos relativos entre categorías ni refleja que, entre
    2017 y 2025, unas divisiones (alimentos, servicios públicos) subieron más o menos que el total (L33).
  * Base = promedio simple de los 12 meses de 2025 (``ANIO_BASE``). El código sigue admitiendo una serie por división (``modo="division"``)
    por si se consigue más adelante; con la serie actual solo existe la división ``00`` (total) y se usa ``modo="total"``.
  * Periodo de referencia de los valores de la ENPH (``referencia``), de la cual sale el denominador del factor:
      - ``mes_recoleccion`` (por defecto): el índice del mes de recolección de cada hogar (``PERIODO`` de la ENPH, AAAAMM). Los hogares
        de Bogotá se levantaron entre 2016-07 y 2017-06. Conserva la inflación dentro del año de campo (factor entre ≈1,56 y ≈1,62
        con la serie de sep-2026) y su promedio casi coincide con el factor de ``promedio_campo``.
      - ``promedio_campo``: un solo factor, con el promedio del índice sobre los meses de recolección observados.
      - ``anio_2017``: un solo factor, con el promedio de 2017 (ene-dic). Subestima la actualización (≈ −2 %) frente a las otras dos porque el
        campo cae sobre todo en 2016-17 (segundo semestre de 2016 + primer semestre de 2017), no en el año calendario 2017.
    ``diag["referencias_equivalentes"]`` reporta el factor que daría cada opción, para la sensibilidad.
  * Supuesto no verificado: los valores de la ENPH están en pesos corrientes del mes de recolección (no se comprobó si el DANE trae
    una variable de actualización).

Entrada del generador: ``outputs/ipc/ipc_total_nacional.csv`` (formato largo, UTF-8, coma): ``mes`` (``AAAA-MM``), ``division`` (``00`` =
total, ``01``…``12`` = divisiones COICOP), ``indice``. Lo produce ``tools/convertir_ipc.py`` desde el xlsx que la autora deja en
``data/raw/IPC/`` (este código NO descarga nada y no escribe en ``data/raw``). Sin el CSV, ``cargar_serie`` devuelve ``None`` y el
generador marca la salida como «sin IPC» (factor 1,0, pesos de 2016-17): nunca se inventa un factor.

Limitaciones declaradas: el IPC mide precios, no corrige crecimiento real del ingreso ni cambios de patrón de consumo desde 2017; la
muestra queda en pesos de PROMEDIO 2025, no de los meses de 2026 de su ventana (a ago-2026 el nivel de precios ya es ≈7 % mayor); la
división ENPH de un artículo es el prefijo de 2 dígitos de su código COICOP (solo relevante si hay serie por división).
"""
from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
import pandas as pd

from sdg import carga

RUTA_IPC = carga.RAIZ / "outputs" / "ipc" / "ipc_total_nacional.csv"
DIR_CRUDO = carga.RAIZ / "data" / "raw" / "IPC"
HOJA_XLSX = "IndicesIPC"
ANIO_BASE = 2025
ANIO_REF_ENPH = 2017
TOTAL = "00"
DIVISIONES = [f"{i:02d}" for i in range(1, 13)]
REFERENCIAS = ("mes_recoleccion", "promedio_campo", "anio_2017")
MESES_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]


class IPCError(ValueError):
    """La serie de IPC no cumple el formato o no cubre los meses que se necesitan."""


def _validar(s: pd.DataFrame, nombre: str) -> pd.DataFrame:
    faltan = {"mes", "division", "indice"} - set(s.columns)
    if faltan:
        raise IPCError(f"{nombre}: faltan columnas {sorted(faltan)}")
    s = s.copy()
    s["division"] = s["division"].astype(str).str.strip().str.zfill(2)
    s["mes"] = s["mes"].astype(str).str.strip()
    s["indice"] = pd.to_numeric(s["indice"], errors="coerce")
    if s["indice"].isna().any() or (s["indice"] <= 0).any():
        raise IPCError(f"{nombre}: índices vacíos o no positivos")
    if not s["mes"].str.fullmatch(r"\d{4}-(0[1-9]|1[0-2])").all():
        raise IPCError(f"{nombre}: la columna mes debe ser AAAA-MM")
    if not s["division"].isin([TOTAL, *DIVISIONES]).all():
        raise IPCError(f"{nombre}: división fuera de 00-12")
    if s.duplicated(["mes", "division"]).any():
        raise IPCError(f"{nombre}: pares (mes, división) repetidos")
    for d, g in s.groupby("division"):
        esperado = pd.period_range(g["mes"].min(), g["mes"].max(), freq="M").astype(str)
        huecos = sorted(set(esperado) - set(g["mes"]))
        if huecos:
            raise IPCError(f"{nombre}: la división {d} tiene meses intermedios sin índice: {huecos[:6]}")
    return s.sort_values(["division", "mes"]).reset_index(drop=True)


def cargar_serie(ruta: Path | str | None = None) -> pd.DataFrame | None:
    """Lee la serie en formato largo y la valida. ``None`` si el archivo no existe (NO es un error: el llamador lo declara)."""
    ruta = Path(ruta) if ruta else RUTA_IPC
    if not ruta.is_file():
        return None
    return _validar(pd.read_csv(ruta, dtype={"mes": str, "division": str}), ruta.name)


def sha256_archivo(ruta: Path | str) -> str:
    return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()


def leer_xlsx_empalme(ruta: Path | str, hoja: str = HOJA_XLSX) -> tuple[pd.DataFrame, dict]:
    """Convierte la hoja de índices del DANE (filas = meses, columnas = años) a formato largo ``mes, division = 00, indice``.

    Estructura que exige: una fila de encabezado con «Mes» en la primera columna y años enteros en las demás, y debajo las 12 filas
    Enero…Diciembre en ese orden. Las celdas vacías solo se admiten al final de la serie (meses aún no publicados); un hueco en medio
    es un error. No se interpreta nada más del archivo (títulos y notas se leen solo para el metadato)."""
    import openpyxl

    ruta = Path(ruta)
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        if hoja not in wb.sheetnames:
            raise IPCError(f"{ruta.name}: no hay una hoja «{hoja}» (hojas: {wb.sheetnames})")
        filas = [tuple(f) for f in wb[hoja].iter_rows(values_only=True)]
    finally:
        wb.close()  # read_only deja el archivo abierto; en Windows bloquea el borrado del temporal
    ih = next((i for i, f in enumerate(filas) if f and isinstance(f[0], str) and f[0].strip() == "Mes"), None)
    if ih is None:
        raise IPCError(f"{ruta.name}/{hoja}: no se encuentra la fila de encabezado con «Mes»")
    encabezado = filas[ih]
    anios = {}
    for j, v in enumerate(encabezado[1:], start=1):
        if v is None:
            continue
        if isinstance(v, bool) or not isinstance(v, (int, float)) or float(v) != int(v) or not 1950 < int(v) < 2100:
            raise IPCError(f"{ruta.name}/{hoja}: el encabezado de la columna {j + 1} no es un año: {v!r}")
        anios[j] = int(v)
    if len(set(anios.values())) != len(anios) or not anios:
        raise IPCError(f"{ruta.name}/{hoja}: años repetidos o ausentes en el encabezado")
    bloque = filas[ih + 1: ih + 13]
    nombres = [str(f[0]).strip() if f and f[0] is not None else "" for f in bloque]
    if nombres != MESES_ES:
        raise IPCError(f"{ruta.name}/{hoja}: las 12 filas bajo el encabezado deben ser Enero…Diciembre; se leyó {nombres}")
    larga = []
    for m, f in enumerate(bloque, start=1):
        for j, anio in anios.items():
            v = f[j] if j < len(f) else None
            if v is None:
                continue
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                raise IPCError(f"{ruta.name}/{hoja}: celda no numérica en {MESES_ES[m - 1]} {anio}: {v!r}")
            larga.append((f"{anio}-{m:02d}", TOTAL, float(v)))
    serie = _validar(pd.DataFrame(larga, columns=["mes", "division", "indice"]), ruta.name)
    textos = [str(f[0]).strip() for f in filas if f and isinstance(f[0], str)]
    dic18 = serie.loc[serie["mes"] == "2018-12", "indice"]
    meta = {
        "archivo": ruta.name, "hoja": hoja, "sha256_xlsx": sha256_archivo(ruta),
        "primer_mes": str(serie["mes"].iloc[0]), "ultimo_mes": str(serie["mes"].iloc[-1]), "n_meses": int(len(serie)),
        "titulo": next((t.replace("\n", " ") for t in textos if "empalme" in t.lower()), None),
        "actualizado": next((t for t in textos if t.lower().startswith("actualizado")), None),
        "fuente": next((t for t in textos if t.lower().startswith("fuente")), None),
        "indice_2018_12": float(dic18.iloc[0]) if len(dic18) else None,
        "base_dic_2018_igual_100": bool(len(dic18) and abs(float(dic18.iloc[0]) - 100.0) < 1e-9),
    }
    return serie, meta


def indice_base(serie: pd.DataFrame, anio: int = ANIO_BASE) -> pd.Series:
    """Promedio simple de los 12 índices mensuales del año, por división. Exige los 12 meses."""
    s = serie[serie["mes"].str.startswith(f"{anio}-")]
    n = s.groupby("division").size()
    if n.empty or (n != 12).any():
        raise IPCError(f"el año base {anio} no tiene los 12 meses en las divisiones {sorted(n[n != 12].index) if len(n) else 'todas'}")
    return s.groupby("division")["indice"].mean()


def _referencias_equivalentes(ancho: pd.DataFrame, base: pd.Series, meses_campo: list) -> dict:
    """Factor total (``base / índice``) que daría cada opción de referencia; solo con la división 00."""
    if TOTAL not in ancho.columns:
        return {}
    b = float(base[TOTAL])
    x = ancho.loc[meses_campo, TOTAL]
    out = {"mes_recoleccion": {"factor_min": float(b / x.max()), "factor_max": float(b / x.min()), "factor_medio_de_meses": float((b / x).mean())},
           "promedio_campo": float(b / x.mean()), "meses_campo": [meses_campo[0], meses_campo[-1], len(meses_campo)]}
    m17 = [f"{ANIO_REF_ENPH}-{m:02d}" for m in range(1, 13)]
    if set(m17) <= set(ancho.index):
        out["anio_2017"] = float(b / ancho.loc[m17, TOTAL].mean())
    return out


def factores(serie: pd.DataFrame | None, periodo: pd.Series, division: pd.Series, modo: str = "division",
             referencia: str = "mes_recoleccion") -> tuple[pd.Series, dict]:
    """Factor de actualización por fila. ``periodo``: ``AAAAMM`` de recolección ENPH. ``division``: 2 dígitos COICOP (otro → total).

    ``modo="division"`` usa la división de cada artículo (si la serie no la trae, cae al total y lo cuenta); ``modo="total"`` usa
    siempre el índice total. ``referencia``: ver la documentación del módulo. Sin serie (``None``) → factor 1,0 y
    ``diag['aplicado'] = False``.
    """
    if modo not in ("division", "total"):
        raise ValueError(f"modo IPC inválido: {modo!r}")
    if referencia not in REFERENCIAS:
        raise ValueError(f"referencia IPC inválida: {referencia!r} (opciones: {REFERENCIAS})")
    if serie is None:
        return pd.Series(1.0, index=periodo.index), {"aplicado": False, "motivo": "no hay serie de IPC (outputs/ipc/ipc_total_nacional.csv)", "modo": modo,
                                                     "referencia": referencia, "pesos": "2016-17 (sin actualizar)"}
    base = indice_base(serie)
    ancho = serie.pivot(index="mes", columns="division", values="indice")
    mes = periodo.str[:4] + "-" + periodo.str[4:6]
    div = pd.Series(TOTAL, index=periodo.index) if modo == "total" else division.where(division.isin(ancho.columns), TOTAL)
    meses_campo = sorted(set(mes))
    faltan = sorted(set(meses_campo) - set(ancho.index))
    if faltan:
        raise IPCError(f"la serie no cubre los meses de recolección {faltan}")
    if referencia == "mes_recoleccion":
        den = ancho.to_numpy()[ancho.index.get_indexer(mes), ancho.columns.get_indexer(div)]
    elif referencia == "promedio_campo":
        den = ancho.loc[meses_campo].mean().reindex(div.to_numpy()).to_numpy()
    else:
        m17 = [f"{ANIO_REF_ENPH}-{m:02d}" for m in range(1, 13)]
        if not set(m17) <= set(ancho.index):
            raise IPCError(f"la serie no trae los 12 meses de {ANIO_REF_ENPH} (referencia anio_2017)")
        den = ancho.loc[m17].mean().reindex(div.to_numpy()).to_numpy()
    den = np.asarray(den, dtype=float)
    if np.isnan(den).any():
        raise IPCError("hay (mes, división) sin índice en la serie")
    f = base.reindex(div.to_numpy()).to_numpy() / den
    n_caida = int((~division.isin(ancho.columns)).sum()) if modo == "division" else 0
    diag = {"aplicado": True, "modo": modo, "referencia": referencia, "anio_base": ANIO_BASE, "base": "promedio simple de los 12 meses de 2025",
            "indice": "total nacional (serie de empalme DANE)" if list(ancho.columns) == [TOTAL] else f"divisiones {list(ancho.columns)}",
            "serie": {"primer_mes": str(ancho.index[0]), "ultimo_mes": str(ancho.index[-1])}, "meses_recoleccion": [meses_campo[0], meses_campo[-1], len(meses_campo)],
            "filas": int(len(f)), "filas_con_caida_al_total": n_caida, "factor_min": float(np.min(f)), "factor_max": float(np.max(f)),
            "indice_base": {k: float(v) for k, v in base.items()}, "referencias_equivalentes": _referencias_equivalentes(ancho, base, meses_campo)}
    return pd.Series(f, index=periodo.index), diag
